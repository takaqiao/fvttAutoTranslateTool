import argparse
import csv
import json
import re
from pathlib import Path
from typing import Any


EN_ALIASES = {
    "english", "en", "source", "src", "term", "key",
    "original", "source_text", "english_term", "source_term",
    "\u82f1\u6587", "\u82f1\u8bed", "\u539f\u6587", "\u672f\u8bed",
}

ZH_ALIASES = {
    "chinese", "zh", "cn", "translation", "translated", "target",
    "preferred_zh", "zh_cn", "target_text", "translated_text",
    "\u4e2d\u6587", "\u7b80\u4f53\u4e2d\u6587", "\u8bd1\u6587", "\u7ffb\u8bd1",
}


def normalize_space(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").strip())


def normalize_header(text: str) -> str:
    return re.sub(r"[\s_\-]+", "", (text or "").strip().casefold())


def contains_en(text: str) -> bool:
    return bool(re.search(r"[A-Za-z]", text or ""))


def parse_delimiter(path: Path, explicit: str) -> str:
    if explicit:
        return explicit
    if path.suffix.lower() in {".tsv", ".tab"}:
        return "\t"
    return ","


def read_text_table(path: Path, encoding: str, delimiter: str) -> tuple[list[str], list[list[str]]]:
    encodings = [encoding]
    if encoding.lower() != "utf-8":
        encodings.append("utf-8")
    if encoding.lower() != "utf-8-sig":
        encodings.append("utf-8-sig")

    last_error: Exception | None = None
    for enc in encodings:
        try:
            with path.open("r", encoding=enc, newline="") as f:
                rows = list(csv.reader(f, delimiter=delimiter))
            break
        except Exception as exc:
            last_error = exc
            rows = []
    else:
        raise RuntimeError(f"Unable to read table file: {path} | {last_error}")

    if not rows:
        return [], []

    headers = [str(x) for x in rows[0]]
    data_rows = [list(map(str, row)) for row in rows[1:]]
    return headers, data_rows


def read_xlsx_table(path: Path, sheet: str, header_row: int) -> tuple[list[str], list[list[str]]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError("Reading .xlsx requires openpyxl. Install with: pip install openpyxl") from exc

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.active

    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else str(v) for v in row])

    if not rows:
        return [], []

    header_idx = max(0, header_row - 1)
    if header_idx >= len(rows):
        raise ValueError(f"header_row={header_row} is out of range")

    headers = rows[header_idx]
    data_rows = rows[header_idx + 1 :]
    return headers, data_rows


def build_header_map(headers: list[str]) -> dict[str, int]:
    out: dict[str, int] = {}
    for idx, h in enumerate(headers):
        key = normalize_header(h)
        if key and key not in out:
            out[key] = idx
    return out


def resolve_column(
    selector: str,
    headers: list[str],
    header_map: dict[str, int],
    aliases: set[str],
    fallback_index: int,
) -> int:
    if selector:
        if selector.isdigit():
            idx = int(selector) - 1
            if idx < 0:
                raise ValueError(f"Column index must be >= 1, got: {selector}")
            return idx

        sel_key = normalize_header(selector)
        if sel_key in header_map:
            return header_map[sel_key]

        raise ValueError(f"Column not found: {selector}")

    for alias in aliases:
        alias_key = normalize_header(alias)
        if alias_key in header_map:
            return header_map[alias_key]

    if len(headers) >= 2:
        return fallback_index

    raise ValueError("Could not auto-detect columns; please provide --en-col and --zh-col")


def safe_cell(row: list[str], idx: int) -> str:
    return row[idx] if 0 <= idx < len(row) else ""


def upsert_glossary(
    glossary: dict[str, str | list[str]],
    key_to_real: dict[str, str],
    english: str,
    chinese: str,
    conflict_mode: str,
) -> tuple[bool, bool]:
    """
    Returns: (added_or_updated, conflict_detected)
    """
    lk = english.casefold()
    existing_key = key_to_real.get(lk)

    if existing_key is None:
        glossary[english] = chinese
        key_to_real[lk] = english
        return True, False

    current = glossary[existing_key]
    if isinstance(current, list):
        current_values = [normalize_space(x) for x in current if normalize_space(x)]
    else:
        current_values = [normalize_space(str(current))] if normalize_space(str(current)) else []

    if chinese in current_values:
        return False, False

    conflict = True
    if conflict_mode == "first":
        return False, conflict

    if conflict_mode == "last":
        glossary[existing_key] = chinese
        return True, conflict

    if conflict_mode == "list":
        merged = current_values + [chinese]
        dedup: list[str] = []
        seen = set()
        for item in merged:
            k = item.casefold()
            if k in seen:
                continue
            seen.add(k)
            dedup.append(item)
        glossary[existing_key] = dedup if len(dedup) > 1 else dedup[0]
        return True, conflict

    raise ValueError(f"Unsupported conflict mode: {conflict_mode}")


def build_terms_payload(
    glossary: dict[str, str | list[str]],
    source_path: Path,
) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    for english in sorted(glossary.keys(), key=lambda s: s.casefold()):
        value = glossary[english]
        if isinstance(value, list):
            variants = {v: 1 for v in value if normalize_space(v)}
            preferred = value[0] if value else ""
        else:
            preferred = normalize_space(str(value))
            variants = {preferred: 1} if preferred else {}

        rows.append(
            {
                "english": english,
                "preferred_zh": preferred,
                "zh_variants": variants,
                "sample_paths": [f"table:{source_path.name}"],
                "total_occurrences": len(variants),
            }
        )

    return {
        "meta": {
            "source_file": str(source_path),
            "total_terms": len(rows),
            "note": "Generated from table. You can edit preferred_zh and reuse in term_governor.",
        },
        "terms": rows,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert table (CSV/TSV/XLSX) to glossary JSON for SF2E translation")
    parser.add_argument("--input", required=True, help="Input table path (.csv/.tsv/.xlsx)")
    parser.add_argument("--output", default="glossary_sf2e_from_table.json", help="Output JSON path")
    parser.add_argument("--format", choices=["dict", "terms"], default="dict", help="Output format")
    parser.add_argument("--en-col", default="", help="English column (header name or 1-based index)")
    parser.add_argument("--zh-col", default="", help="Chinese column (header name or 1-based index)")
    parser.add_argument("--sheet", default="", help="XLSX sheet name (default: active)")
    parser.add_argument("--header-row", type=int, default=1, help="XLSX header row index, 1-based")
    parser.add_argument("--encoding", default="utf-8-sig", help="Text table encoding")
    parser.add_argument("--delimiter", default="", help="CSV delimiter, default auto by extension")
    parser.add_argument("--conflict-mode", choices=["first", "last", "list"], default="last", help="Duplicate EN key strategy")
    parser.add_argument("--merge-existing", default="", help="Optional existing glossary JSON to merge into")
    parser.add_argument("--conflicts-out", default="", help="Optional conflict report JSON path")
    args = parser.parse_args()

    in_path = Path(args.input)
    out_path = Path(args.output)
    if not in_path.exists():
        raise FileNotFoundError(f"Input file not found: {in_path}")

    suffix = in_path.suffix.lower()
    if suffix in {".xlsx"}:
        headers, data_rows = read_xlsx_table(in_path, args.sheet, args.header_row)
    elif suffix in {".csv", ".tsv", ".tab"}:
        delimiter = parse_delimiter(in_path, args.delimiter)
        headers, data_rows = read_text_table(in_path, args.encoding, delimiter)
    else:
        raise ValueError(f"Unsupported input extension: {suffix}")

    if not headers:
        raise ValueError("No header row detected")

    header_map = build_header_map(headers)
    en_col = resolve_column(args.en_col, headers, header_map, EN_ALIASES, fallback_index=0)
    zh_col = resolve_column(args.zh_col, headers, header_map, ZH_ALIASES, fallback_index=1)

    glossary: dict[str, str | list[str]] = {}
    key_to_real: dict[str, str] = {}

    if args.merge_existing:
        merge_path = Path(args.merge_existing)
        if merge_path.exists():
            existing = json.loads(merge_path.read_text(encoding="utf-8"))
            if isinstance(existing, dict):
                for k, v in existing.items():
                    kk = normalize_space(str(k))
                    if not kk:
                        continue
                    glossary[kk] = v
                    key_to_real[kk.casefold()] = kk

    row_total = 0
    accepted = 0
    skipped_empty = 0
    skipped_non_en = 0
    conflicts: list[dict[str, Any]] = []

    for row_idx, row in enumerate(data_rows, start=2):
        row_total += 1
        en = normalize_space(safe_cell(row, en_col))
        zh = normalize_space(safe_cell(row, zh_col))

        if not en or not zh:
            skipped_empty += 1
            continue

        if not contains_en(en):
            skipped_non_en += 1
            continue

        changed, conflict = upsert_glossary(glossary, key_to_real, en, zh, args.conflict_mode)
        if changed:
            accepted += 1
        if conflict:
            conflicts.append(
                {
                    "row": row_idx,
                    "english": en,
                    "incoming_zh": zh,
                    "current_value": glossary[key_to_real[en.casefold()]],
                }
            )

    if args.format == "terms":
        payload: Any = build_terms_payload(glossary, in_path)
    else:
        payload = dict(sorted(glossary.items(), key=lambda item: item[0].casefold()))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    if args.conflicts_out:
        conflict_path = Path(args.conflicts_out)
        conflict_path.parent.mkdir(parents=True, exist_ok=True)
        conflict_path.write_text(json.dumps(conflicts, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Input: {in_path}")
    print(f"Output: {out_path}")
    print(f"Rows: {row_total} | Accepted: {accepted} | Empty skipped: {skipped_empty} | Non-EN skipped: {skipped_non_en}")
    print(f"Glossary entries: {len(glossary)} | Conflicts: {len(conflicts)} | Conflict mode: {args.conflict_mode}")


if __name__ == "__main__":
    main()
