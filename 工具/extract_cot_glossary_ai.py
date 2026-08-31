"""
COT 双语 PDF 交叉术语提取工具（AI 版）

读取英文原版 PDF + 中文翻译 PDF/docx/txt，
按页对齐后送 AI 交叉比对，提取完整的中英术语对。
可选参照 glossary.json 提供上下文。

用法:
  python extract_cot_glossary_ai.py [--config glossary_extract_config_cot.json]
"""

import argparse
import concurrent.futures
import hashlib
import json
import os
import re
import sys
import io
import time
from pathlib import Path
from threading import Lock

from openai import OpenAI
from tqdm import tqdm

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx as python_docx
except ImportError:
    python_docx = None

# Windows 控制台编码修复
if sys.stdout.encoding and sys.stdout.encoding.lower().startswith("gbk"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

CONFIG_PATH = Path("glossary_extract_config_cot.json")

_log_lock = Lock()
_cache_lock = Lock()
_stats_lock = Lock()


# ---------------------------------------------------------------------------
# 工具函数
# ---------------------------------------------------------------------------

def _write_log(log_path: Path, msg: str):
    ts = time.strftime("%H:%M:%S", time.localtime())
    with _log_lock:
        try:
            with log_path.open("a", encoding="utf-8") as f:
                f.write(f"[{ts}] {msg}\n")
        except Exception:
            pass


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "")).strip()


def _hash_text(text: str) -> str:
    return hashlib.sha1((text or "").encode("utf-8", errors="ignore")).hexdigest()


def _contains_zh(text: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", text or ""))


def _contains_en(text: str) -> bool:
    return bool(re.search(r"[A-Za-z]{2,}", text or ""))


# ---------------------------------------------------------------------------
# 文件读取
# ---------------------------------------------------------------------------

def read_pdf_pages(pdf_path: Path, page_start: int = 1, page_end: int | None = None) -> list[tuple[int, str]]:
    """读取 PDF 返回 [(page_no, text), ...]"""
    if fitz is None:
        raise RuntimeError("需要 PyMuPDF: pip install pymupdf")
    pages = []
    with fitz.open(str(pdf_path)) as doc:
        total = len(doc)
        start = max(1, page_start)
        end = total if page_end is None else min(total, page_end)
        for pn in range(start, end + 1):
            text = doc[pn - 1].get_text("text") or ""
            text = text.replace("\x00", "")
            pages.append((pn, text))
    return pages


def read_txt(path: Path) -> str:
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return ""


def read_docx(path: Path) -> str:
    if python_docx is None:
        print(f"  ⚠ 需要 python-docx 来读取 docx: pip install python-docx")
        return ""
    try:
        doc = python_docx.Document(str(path))
        parts = []
        for para in doc.paragraphs:
            parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    except Exception as e:
        print(f"  ⚠ 读取 docx 失败: {path.name} | {e}")
        return ""


def read_extra_file(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".txt":
        return read_txt(path)
    elif ext == ".docx":
        return read_docx(path)
    elif ext == ".pdf":
        pages = read_pdf_pages(path)
        return "\n".join(t for _, t in pages)
    return ""


# ---------------------------------------------------------------------------
# 配置加载
# ---------------------------------------------------------------------------

def load_config(config_path: Path) -> dict:
    if not config_path.exists():
        print(f"❌ 配置文件不存在: {config_path}")
        sys.exit(1)
    try:
        cfg = json.loads(config_path.read_text(encoding="utf-8"))
        if not isinstance(cfg, dict):
            raise ValueError("配置文件应为 JSON 对象")
        return cfg
    except Exception as e:
        print(f"❌ 配置文件读取失败: {e}")
        sys.exit(1)


# ---------------------------------------------------------------------------
# 批次构建
# ---------------------------------------------------------------------------

def build_aligned_batches(
    en_pages: list[tuple[int, str]],
    zh_pages: list[tuple[int, str]],
    pages_per_batch: int,
    max_chars: int,
    page_offset: int = 0,
) -> list[dict]:
    """将英文和中文页面对齐分批。

    page_offset: 中文页码 = 英文页码 + offset（用于补偿页码偏移）
    """
    en_dict = {pn: text for pn, text in en_pages}
    zh_dict = {pn: text for pn, text in zh_pages}

    # 所有涉及的页码范围
    all_en_pages = sorted(en_dict.keys())
    if not all_en_pages:
        return []

    batches = []
    i = 0
    while i < len(all_en_pages):
        batch_en_parts = []
        batch_zh_parts = []
        batch_en_pns = []
        batch_zh_pns = []
        chars = 0

        while i < len(all_en_pages):
            en_pn = all_en_pages[i]
            zh_pn = en_pn + page_offset

            en_text = en_dict.get(en_pn, "")
            zh_text = zh_dict.get(zh_pn, "")

            part_len = len(en_text) + len(zh_text) + 60  # 标记开销

            if batch_en_pns and (len(batch_en_pns) >= pages_per_batch or chars + part_len > max_chars):
                break

            if en_text.strip():
                batch_en_parts.append(f"=== EN Page {en_pn} ===\n{en_text}")
                batch_en_pns.append(en_pn)
            if zh_text.strip():
                batch_zh_parts.append(f"=== ZH Page {zh_pn} ===\n{zh_text}")
                batch_zh_pns.append(zh_pn)
            chars += part_len
            i += 1

        if batch_en_parts or batch_zh_parts:
            batches.append({
                "en_text": "\n\n".join(batch_en_parts),
                "zh_text": "\n\n".join(batch_zh_parts),
                "en_pages": batch_en_pns,
                "zh_pages": batch_zh_pns,
                "label": f"EN p{batch_en_pns[0]}-{batch_en_pns[-1]}" if batch_en_pns else "extra",
            })

    return batches


def build_extra_batches(extra_texts: list[tuple[str, str]], max_chars: int) -> list[dict]:
    """将额外的中文文件（txt/docx）分批。"""
    batches = []
    for filename, text in extra_texts:
        if not text.strip():
            continue
        # 按 max_chars 切分
        chunks = []
        start = 0
        while start < len(text):
            end = start + max_chars
            chunks.append(text[start:end])
            start = end
        for ci, chunk in enumerate(chunks):
            batches.append({
                "en_text": "",
                "zh_text": f"=== {filename} (part {ci+1}) ===\n{chunk}",
                "en_pages": [],
                "zh_pages": [],
                "label": f"{filename} part{ci+1}",
            })
    return batches


# ---------------------------------------------------------------------------
# AI 提取
# ---------------------------------------------------------------------------

SYSTEM_PROMPT = """\
你是 PF2E/SF2E TRPG 术语提取专家。你会收到英文原文和对应的中文翻译文本。

你的任务：
1. 从英文原文中识别所有 **专有名词、游戏术语、地名、人名、怪物名、法术名、专长名、物品名、\
技能名、动作名、特征名、变体名、组织名、章节名** 等术语。
2. 从中文翻译中找到这些术语的对应翻译。
3. 同时提取中文文本中以 中文（English）格式明确标注的术语对。
4. 每个术语只需输出一次，选最准确的翻译。

输出要求：
- 严格 JSON 数组，每项格式：{"english":"...","chinese":"...","category":"..."}
- category 可选值：creature, spell, feat, item, action, skill, location, npc, organization, \
chapter, trait, archetype, hazard, artifact, term
- 不要猜测或编造不存在于文本中的术语
- 英文保持原文大小写
- 中文只输出术语名，不要输出整句话
- 若无可提取术语，输出 []

示例输出：
[
  {"english":"Seldeg Bhedlis","chinese":"塞尔德格·贝德利斯","category":"npc"},
  {"english":"Gravelands","chinese":"死坟末土","category":"location"},
  {"english":"Shield Block","chinese":"盾牌格挡","category":"action"}
]\
"""

SYSTEM_PROMPT_ZH_ONLY = """\
你是 PF2E/SF2E TRPG 术语提取专家。你会收到中文翻译文本（可能包含括号标注的英文术语）。

你的任务：
1. 提取文本中所有以 中文（English）或 中文 English 格式出现的术语对。
2. 识别所有专有名词（人名、地名、怪物名、法术名、专长名、物品名等）。
3. 对于只有中文没有英文的专有名词，若你确信知道其英文原文，也可以输出。

输出要求：
- 严格 JSON 数组，每项格式：{"english":"...","chinese":"...","category":"..."}
- category 可选值：creature, spell, feat, item, action, skill, location, npc, organization, \
chapter, trait, archetype, hazard, artifact, term
- 中文只输出术语名，不要输出整句话
- 若无可提取术语，输出 []
"""


class TermExtractorAI:
    def __init__(self, cfg: dict):
        api_key = cfg.get("openai_api_key") or os.getenv("OPENAI_API_KEY", "")
        if not api_key:
            raise RuntimeError("未设置 openai_api_key 且环境变量 OPENAI_API_KEY 为空")
        self.client = OpenAI(
            api_key=api_key,
            base_url=cfg.get("openai_base_url", "https://api.openai.com/v1"),
        )
        self.model = cfg.get("model", "gpt-4.1")
        self.max_retries = max(1, int(cfg.get("max_retries", 3)))

    def extract(self, batch: dict, ref_terms_text: str = "") -> list[dict]:
        en_text = batch.get("en_text", "")
        zh_text = batch.get("zh_text", "")

        has_en = bool(en_text.strip())
        has_zh = bool(zh_text.strip())

        if has_en and has_zh:
            instructions = SYSTEM_PROMPT
            user_input = (
                "[英文原文]\n" + en_text +
                "\n\n[中文翻译]\n" + zh_text
            )
        elif has_zh:
            instructions = SYSTEM_PROMPT_ZH_ONLY
            user_input = "[中文文本]\n" + zh_text
        else:
            return []

        if ref_terms_text:
            user_input += "\n\n[参考术语表（已有翻译，优先保持一致）]\n" + ref_terms_text

        user_input += "\n\n请提取所有术语，返回 JSON 数组。"

        for attempt in range(1, self.max_retries + 1):
            try:
                resp = self.client.responses.create(
                    model=self.model,
                    instructions=instructions,
                    input=user_input,
                )
                raw = (resp.output_text or "").strip()
                return self._parse_response(raw)
            except Exception as e:
                if attempt == self.max_retries:
                    raise RuntimeError(f"AI 提取失败: {batch.get('label', '?')} | {e}") from e
                time.sleep(attempt * 2)
        return []

    def _parse_response(self, raw: str) -> list[dict]:
        if not raw:
            return []
        # 提取 JSON 数组
        fence = re.search(r"```(?:json)?\s*(\[[\s\S]*?\])\s*```", raw, re.I)
        if fence:
            text = fence.group(1)
        else:
            start = raw.find("[")
            end = raw.rfind("]")
            if start != -1 and end > start:
                text = raw[start:end + 1]
            else:
                return []
        try:
            arr = json.loads(text)
            if not isinstance(arr, list):
                return []
        except json.JSONDecodeError:
            return []

        results = []
        for item in arr:
            if not isinstance(item, dict):
                continue
            en = _norm(str(item.get("english", "")))
            zh = _norm(str(item.get("chinese", "")))
            cat = str(item.get("category", "term")).strip().lower()
            if not en or not zh or not _contains_en(en) or not _contains_zh(zh):
                continue
            if len(zh) > 40 or len(en) > 100:
                continue
            results.append({"english": en, "chinese": zh, "category": cat})
        return results


# ---------------------------------------------------------------------------
# 限速
# ---------------------------------------------------------------------------

class RateLimiter:
    def __init__(self, rpm: int):
        self.rpm = max(1, rpm)
        self.min_interval = 60.0 / self.rpm
        self.last_request = 0.0
        self.lock = Lock()

    def wait(self):
        with self.lock:
            now = time.time()
            gap = self.last_request + self.min_interval - now
            if gap > 0:
                time.sleep(gap)
            self.last_request = time.time()


# ---------------------------------------------------------------------------
# 缓存
# ---------------------------------------------------------------------------

def _load_cache(path: Path, enabled: bool) -> dict:
    if not enabled or not path.exists():
        return {"batches": {}}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, dict) and isinstance(data.get("batches"), dict):
            return data
    except Exception:
        pass
    return {"batches": {}}


def _save_cache(path: Path, cache: dict, enabled: bool):
    if not enabled:
        return
    try:
        path.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass


# ---------------------------------------------------------------------------
# 合并 & 输出
# ---------------------------------------------------------------------------

def merge_all_terms(all_terms: list[dict]) -> list[dict]:
    """按英文 key 合并去重，保留出现次数最多的中文翻译。"""
    bucket: dict[str, dict] = {}
    for item in all_terms:
        key = item["english"].lower().strip()
        if key not in bucket:
            bucket[key] = {
                "english": item["english"],
                "zh_variants": {},
                "categories": {},
            }
        entry = bucket[key]
        zh = item["chinese"]
        cat = item.get("category", "term")
        entry["zh_variants"][zh] = entry["zh_variants"].get(zh, 0) + 1
        entry["categories"][cat] = entry["categories"].get(cat, 0) + 1
        if len(item["english"]) > len(entry["english"]):
            entry["english"] = item["english"]

    results = []
    for entry in sorted(bucket.values(), key=lambda x: x["english"].lower()):
        sorted_zh = sorted(entry["zh_variants"].items(), key=lambda x: x[1], reverse=True)
        sorted_cat = sorted(entry["categories"].items(), key=lambda x: x[1], reverse=True)
        results.append({
            "english": entry["english"],
            "chinese": sorted_zh[0][0],
            "category": sorted_cat[0][0],
            "occurrences": sum(v for v in entry["zh_variants"].values()),
            "zh_variants": {k: v for k, v in sorted_zh} if len(sorted_zh) > 1 else {},
        })
    return results


def cross_reference_glossary(terms: list[dict], glossary: dict) -> list[dict]:
    """标注是否已在参考术语表中，以及译名差异。"""
    lower_glossary = {}
    for en, zh in glossary.items():
        lower_glossary[en.lower().strip()] = {"en": en, "zh": zh}

    for term in terms:
        key = term["english"].lower().strip()
        if key in lower_glossary:
            existing = lower_glossary[key]
            term["in_glossary"] = True
            term["glossary_zh"] = existing["zh"] if isinstance(existing["zh"], str) else ", ".join(existing["zh"])
            term["zh_differs"] = term["chinese"] != term["glossary_zh"]
        else:
            term["in_glossary"] = False
            term["glossary_zh"] = ""
            term["zh_differs"] = False
    return terms


def build_ref_terms_text(glossary: dict, en_text: str, max_terms: int = 50) -> str:
    """从参考术语表中挑选与当前批次内容相关的术语，作为上下文。"""
    if not glossary or not en_text:
        return ""
    en_lower = en_text.lower()
    relevant = []
    for en, zh in glossary.items():
        if en.lower() in en_lower:
            zh_str = zh if isinstance(zh, str) else ", ".join(zh)
            relevant.append(f"{en}: {zh_str}")
    relevant = relevant[:max_terms]
    return "\n".join(relevant) if relevant else ""


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="COT 双语 PDF 交叉术语提取 (AI)")
    parser.add_argument("--config", default=str(CONFIG_PATH), help="配置文件路径")
    args = parser.parse_args()

    cfg = load_config(Path(args.config))
    log_path = Path(cfg.get("log_path", "glossary_extract_cot.log"))

    if fitz is None:
        print("❌ 需要 PyMuPDF: pip install pymupdf")
        return

    # ── 读取 PDF ──
    en_pdf = Path(cfg.get("en_pdf", ""))
    zh_pdf = Path(cfg.get("zh_pdf", ""))

    if not en_pdf.exists():
        print(f"❌ 英文 PDF 不存在: {en_pdf}")
        return
    if not zh_pdf.exists():
        print(f"❌ 中文 PDF 不存在: {zh_pdf}")
        return

    en_start = int(cfg.get("en_page_start", 1))
    en_end = cfg.get("en_page_end")
    zh_start = int(cfg.get("zh_page_start", 1))
    zh_end = cfg.get("zh_page_end")
    page_offset = int(cfg.get("page_offset", 0))

    print(f"读取英文 PDF: {en_pdf}")
    en_pages = read_pdf_pages(en_pdf, en_start, en_end)
    print(f"  {len(en_pages)} 页")

    print(f"读取中文 PDF: {zh_pdf}")
    zh_pages = read_pdf_pages(zh_pdf, zh_start, zh_end)
    print(f"  {len(zh_pages)} 页")

    # ── 构建对齐批次 ──
    pages_per_batch = int(cfg.get("pages_per_batch", 6))
    max_chars = int(cfg.get("max_chars_per_batch", 24000))

    batches = build_aligned_batches(en_pages, zh_pages, pages_per_batch, max_chars, page_offset)
    print(f"对齐批次: {len(batches)}")

    # ── 额外中文文件 ──
    extra_files = cfg.get("zh_extra_files", [])
    extra_texts = []
    for fp in extra_files:
        p = Path(fp)
        if p.exists():
            text = read_extra_file(p)
            if text.strip():
                extra_texts.append((p.name, text))
                print(f"  额外文件: {p.name} ({len(text)} 字符)")
    extra_batches = build_extra_batches(extra_texts, max_chars)
    if extra_batches:
        print(f"额外批次: {len(extra_batches)}")
        batches.extend(extra_batches)

    print(f"总批次: {len(batches)}")

    # ── 加载参考术语表 ──
    ref_glossary_path = cfg.get("reference_glossary", "")
    ref_glossary = {}
    if ref_glossary_path:
        p = Path(ref_glossary_path)
        if p.exists():
            try:
                ref_glossary = json.loads(p.read_text(encoding="utf-8-sig"))
                if not isinstance(ref_glossary, dict):
                    ref_glossary = {}
                print(f"参考术语表: {p} ({len(ref_glossary)} 条)")
            except Exception as e:
                print(f"⚠ 参考术语表读取失败: {e}")

    # ── 加载已有 COT 术语表（也作为 AI 参考） ──
    existing_glossary_path = cfg.get("existing_glossary", "")
    existing_glossary = {}
    if existing_glossary_path:
        p = Path(existing_glossary_path)
        if p.exists():
            try:
                existing_glossary = json.loads(p.read_text(encoding="utf-8-sig"))
                if not isinstance(existing_glossary, dict):
                    existing_glossary = {}
                print(f"增量基础术语表: {p} ({len(existing_glossary)} 条)")
            except Exception as e:
                print(f"⚠ 增量基础术语表读取失败: {e}")

    # 合并参考术语：ref_glossary + existing_glossary，已有 COT 术语优先
    combined_ref_glossary = dict(ref_glossary)
    combined_ref_glossary.update(existing_glossary)
    if existing_glossary:
        print(f"合并参考术语: {len(combined_ref_glossary)} 条 (通用 {len(ref_glossary)} + COT {len(existing_glossary)})")

    # ── 缓存 ──
    cache_enabled = bool(cfg.get("cache_enabled", True))
    cache_path = Path(cfg.get("cache_path", "glossary_extract_cache_cot.json"))
    cache = _load_cache(cache_path, cache_enabled)

    # ── AI 提取器 ──
    model = cfg.get("model", "gpt-4.1")
    extractor = TermExtractorAI(cfg)
    limiter = RateLimiter(int(cfg.get("target_rpm", 30)))

    max_workers = max(1, int(cfg.get("max_workers", 4)))
    stats = {
        "batch_count": len(batches),
        "cache_hit": 0,
        "ai_ok": 0,
        "ai_fail": 0,
    }

    all_terms = []

    def process_batch(idx: int, batch: dict) -> list[dict]:
        label = batch["label"]
        batch_hash = _hash_text(model + "|" + batch["en_text"] + batch["zh_text"])
        cache_key = f"batch:{idx}:{label}"

        # 检查缓存
        with _cache_lock:
            cached = cache["batches"].get(cache_key)
            if (
                cache_enabled
                and isinstance(cached, dict)
                and cached.get("hash") == batch_hash
                and isinstance(cached.get("terms"), list)
            ):
                with _stats_lock:
                    stats["cache_hit"] += 1
                return cached["terms"]

        # 构建参考术语上下文
        ref_text = build_ref_terms_text(combined_ref_glossary, batch["en_text"] + " " + batch["zh_text"], 80)

        try:
            limiter.wait()
            terms = extractor.extract(batch, ref_text)
            with _stats_lock:
                stats["ai_ok"] += 1

            # 缓存结果
            if cache_enabled:
                with _cache_lock:
                    cache["batches"][cache_key] = {
                        "hash": batch_hash,
                        "terms": terms,
                    }
            return terms
        except Exception as e:
            _write_log(log_path, f"ERROR {label} | {e}")
            with _stats_lock:
                stats["ai_fail"] += 1
            return []

    # ── 执行提取 ──
    print(f"\n开始 AI 术语提取 (model={cfg.get('model')}, workers={max_workers})...")
    _write_log(log_path, f"START batches={len(batches)} model={cfg.get('model')}")

    if max_workers <= 1:
        for idx, batch in enumerate(tqdm(batches, desc="术语提取", dynamic_ncols=True)):
            terms = process_batch(idx, batch)
            all_terms.extend(terms)
    else:
        with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = {executor.submit(process_batch, idx, batch): idx for idx, batch in enumerate(batches)}
            for fut in tqdm(concurrent.futures.as_completed(futures), total=len(futures), desc="术语提取", dynamic_ncols=True):
                terms = fut.result()
                all_terms.extend(terms)

    # 中间保存缓存
    _save_cache(cache_path, cache, cache_enabled)

    print(f"\n原始提取: {len(all_terms)} 个术语对")

    # ── 合并去重 ──
    merged = merge_all_terms(all_terms)
    print(f"去重合并: {len(merged)} 个唯一术语")

    # ── 交叉参照 ──
    if ref_glossary:
        merged = cross_reference_glossary(merged, ref_glossary)

    # ── 输出 ──
    output_json = Path(cfg.get("output_json", "glossary_cot_ai.json"))
    output_candidates = Path(cfg.get("output_candidates_json", "glossary_cot_ai_candidates.json"))
    output_conflicts = Path(cfg.get("output_conflicts_json", "glossary_cot_ai_conflicts.json"))

    # AI 提取的术语表
    ai_glossary = {item["english"]: item["chinese"] for item in merged}

    # 增量合并：以已有术语表为基础，AI 新提取的条目追加，冲突时保留已有
    glossary_dict = dict(existing_glossary)
    ai_added = 0
    ai_conflicts = []
    lower_existing = {k.lower().strip(): k for k in glossary_dict}
    for en, zh in ai_glossary.items():
        lk = en.lower().strip()
        if lk in lower_existing:
            existing_key = lower_existing[lk]
            old_zh = glossary_dict[existing_key]
            if old_zh != zh:
                ai_conflicts.append({"english": en, "existing_zh": old_zh, "ai_zh": zh})
        else:
            glossary_dict[en] = zh
            lower_existing[lk] = en
            ai_added += 1

    if existing_glossary:
        print(f"增量合并: 已有 {len(existing_glossary)} + AI新增 {ai_added} = {len(glossary_dict)} 条 (冲突 {len(ai_conflicts)} 保留已有)")

    output_json.write_text(json.dumps(glossary_dict, ensure_ascii=False, indent=2), encoding="utf-8")

    # 分类统计
    new_terms = [x for x in merged if not x.get("in_glossary", False)]
    existing_terms = [x for x in merged if x.get("in_glossary", False)]
    differing = [x for x in merged if x.get("zh_differs", False)]
    categories = {}
    for item in merged:
        cat = item.get("category", "term")
        categories[cat] = categories.get(cat, 0) + 1

    # 详细报告
    payload = {
        "meta": {
            "adventure": "Claws of the Tyrant / 暴君之爪",
            "ai_extracted_terms": len(merged),
            "new_vs_ref_glossary": len(new_terms),
            "existing_in_ref_glossary": len(existing_terms),
            "translation_differs_vs_ref": len(differing),
            "incremental_base": len(existing_glossary),
            "incremental_added": ai_added,
            "incremental_conflicts": len(ai_conflicts),
            "final_glossary_size": len(glossary_dict),
            "categories": dict(sorted(categories.items())),
            "stats": stats,
            "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        },
        "glossary": glossary_dict,
        "ai_terms_detail": merged,
        "ai_new_terms_only": {x["english"]: x["chinese"] for x in new_terms},
        "differing_vs_ref_glossary": [
            {"english": x["english"], "ai_zh": x["chinese"], "glossary_zh": x["glossary_zh"]}
            for x in differing
        ],
        "incremental_conflicts": ai_conflicts,
    }
    output_candidates.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    # 冲突报告
    conflict_payload = {
        "meta": {
            "vs_ref_glossary": len(differing),
            "vs_existing_cot": len(ai_conflicts),
        },
        "vs_ref_glossary": payload["differing_vs_ref_glossary"],
        "vs_existing_cot": ai_conflicts,
    }
    output_conflicts.write_text(json.dumps(conflict_payload, ensure_ascii=False, indent=2), encoding="utf-8")

    # ── 打印摘要 ──
    print(f"\n{'='*55}")
    print(f"冒险: Claws of the Tyrant / 暴君之爪")
    print(f"AI 提取术语: {len(merged)}")
    print(f"  vs ref glossary: 新={len(new_terms)} 已有={len(existing_terms)} 差异={len(differing)}")
    if existing_glossary:
        print(f"增量合并: 基础={len(existing_glossary)} + AI新增={ai_added} = {len(glossary_dict)} (冲突={len(ai_conflicts)})")
    print(f"分类:")
    for cat, count in sorted(categories.items()):
        print(f"  {cat}: {count}")
    print(f"{'='*55}")
    print(f"统计: 批次={stats['batch_count']} 缓存命中={stats['cache_hit']} "
          f"AI成功={stats['ai_ok']} AI失败={stats['ai_fail']}")
    print(f"\n✅ 术语表: {output_json} ({len(glossary_dict)} 条)")
    print(f"✅ 详细报告: {output_candidates}")
    print(f"✅ 差异报告: {output_conflicts}")

    # 保存缓存
    _save_cache(cache_path, cache, cache_enabled)
    _write_log(log_path, f"DONE terms={len(merged)} new={len(new_terms)} differs={len(differing)}")


if __name__ == "__main__":
    main()
