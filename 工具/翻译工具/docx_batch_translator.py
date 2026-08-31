import concurrent.futures
import hashlib
import json
import os
import re
import time
from dataclasses import dataclass
from pathlib import Path
from threading import Lock

from openai import OpenAI
from tqdm import tqdm

try:
    from docx import Document
except Exception:
    Document = None


CONFIG_PATH = Path("docx_translate_config.json")
DEFAULT_CONFIG = {
    "openai_api_key": "",
    "openai_base_url": "https://api.openai.com/v1",
    "model": "gpt-5.2",
    "fallback_on_failure_enabled": True,
    "fallback_model_on_failure": "gpt-4.1",
    "request_timeout_seconds": 60,
    "fallback_timeout_seconds": 60,
    "input_docx_path": "",
    "input_folder_path": "",
    "input_file_glob": "*.docx",
    "input_folder_recursive": False,
    "output_dir": "docx_output",
    "append_original_text": True,
    "glossary_path": "glossary.json",
    "prompt_profile": "default",
    "max_workers": 8,
    "target_rpm": 300,
    "max_retries": 3,
    "batch_size": 20,
    "cache_enabled": True,
    "cache_path": "docx_translate_cache.json",
    "log_path": "docx_translate_run.log",
    "output_markdown": True,
    "output_json": True,
    "top_n_slowest": 20,
}


_log_lock = Lock()
_cache_lock = Lock()
_stats_lock = Lock()


def _write_log(log_path: Path, msg: str):
    ts = time.strftime("%H:%M:%S", time.localtime())
    line = f"[{ts}] {msg}\n"
    with _log_lock:
        try:
            with log_path.open("a", encoding="utf-8") as f:
                f.write(line)
        except Exception:
            pass


def _load_config() -> dict:
    cfg = dict(DEFAULT_CONFIG)
    if CONFIG_PATH.exists():
        try:
            user_cfg = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
            if isinstance(user_cfg, dict):
                cfg.update(user_cfg)
        except Exception:
            pass
    else:
        try:
            CONFIG_PATH.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
            print(f"已生成默认配置：{CONFIG_PATH}")
        except Exception:
            pass
    return cfg


def _hash_text(text: str) -> str:
    return hashlib.sha1((text or "").encode("utf-8", errors="ignore")).hexdigest()


def _strip_code_like_tokens(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r"@UUID\[[^\]]+\]", " ", text)
    text = re.sub(r"@Compendium\[[^\]]+\]", " ", text)
    text = re.sub(r"@Localize\[[^\]]+\]", " ", text)
    text = re.sub(r"\[\[.*?\]\]", " ", text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"&[a-zA-Z0-9#]+;", " ", text)
    return text


def _contains_english(text: str) -> bool:
    return bool(re.search(r"[A-Za-z]", _strip_code_like_tokens(text or "")))


def _contains_chinese(text: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", _strip_code_like_tokens(text or "")))


def _sanitize_text(text: str) -> str:
    if not text:
        return ""
    return re.sub(r"[ \t\x0b\x0c\r]+", " ", text).replace("\u00a0", " ").strip()


def _parse_json_block(raw: str):
    if not raw:
        return {}
    text = raw.strip()
    fence_match = re.search(r"```(?:json)?\s*(\{[\s\S]*\})\s*```", text, flags=re.IGNORECASE)
    if fence_match:
        text = fence_match.group(1).strip()
    else:
        start = text.find("{")
        end = text.rfind("}")
        if start != -1 and end != -1 and end > start:
            text = text[start:end + 1]
    try:
        obj = json.loads(text)
        return obj if isinstance(obj, dict) else {}
    except Exception:
        return {}


def _load_cache(path: Path, enabled: bool) -> dict:
    if not enabled or not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _save_cache(path: Path, cache: dict, enabled: bool):
    if not enabled:
        return
    try:
        path.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass


def _load_glossary(path: Path):
    if not path.exists():
        return {}, [], 0
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except Exception:
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return {}, [], 0
    if not isinstance(data, dict):
        return {}, [], 0

    buckets = {}
    terms_for_replace = []
    count = 0

    for en_raw, cn_raw in data.items():
        en = str(en_raw).strip()
        if not en:
            continue
        if isinstance(cn_raw, list):
            cns = [str(x).strip() for x in cn_raw if str(x).strip()]
            cn = cns[0] if cns else ""
        else:
            cn = str(cn_raw).strip()
            cns = [cn] if cn else []
        if not cns:
            continue

        letters = re.sub(r"[^A-Za-z]", "", en)
        if letters:
            pattern = r"(?<![A-Za-z])" + r"\s*".join(list(letters)) + r"(?![A-Za-z])"
            first = letters[0].lower()
            length = len(letters)
            length_bucket = (length // 5) * 5
            buckets.setdefault(first, {}).setdefault(length_bucket, []).append((
                length,
                re.compile(pattern, flags=re.IGNORECASE),
                cns,
                en,
            ))

        escaped_en = re.escape(en)
        replace_pattern = re.compile(rf"(?<![A-Za-z]){escaped_en}(?![A-Za-z])", flags=re.IGNORECASE)
        terms_for_replace.append((len(en), replace_pattern, cn))
        count += 1

    for first in list(buckets.keys()):
        for length_bucket in list(buckets[first].keys()):
            buckets[first][length_bucket].sort(key=lambda x: x[0], reverse=True)

    terms_for_replace.sort(key=lambda x: x[0], reverse=True)
    return buckets, terms_for_replace, count


def _find_glossary_matches(text: str, glossary_buckets: dict):
    if not text or not glossary_buckets:
        return []
    letters = {w[0].lower() for w in re.findall(r"[A-Za-z]+", text)}
    if not letters:
        return []

    matched = []
    for first in letters:
        length_map = glossary_buckets.get(first)
        if not length_map:
            continue
        for bucket_terms in length_map.values():
            for _, pattern, cns, en in bucket_terms:
                if pattern.search(text):
                    matched.append((en, cns[0]))

    dedup = {}
    for en, cn in matched:
        dedup[en.lower()] = (en, cn)

    out = list(dedup.values())
    out.sort(key=lambda x: len(x[0]), reverse=True)
    return out


def _apply_glossary_post_translation(text: str, terms_for_replace):
    if not text or not terms_for_replace:
        return text
    out = text
    for _, pattern, cn in terms_for_replace:
        out = pattern.sub(cn, out)
    return out


class RateLimiter:
    def __init__(self, rpm: int):
        self.rpm = max(int(rpm), 1)
        self.capacity = self.rpm
        self.tokens = self.capacity
        self.last_refill = time.time()
        self.min_interval = 60.0 / self.rpm
        self.last_request = 0.0
        self.lock = Lock()

    def wait(self):
        with self.lock:
            while True:
                now = time.time()
                elapsed = now - self.last_refill
                if elapsed >= 60.0:
                    refill = int(elapsed // 60.0) * self.capacity
                    self.tokens = min(self.capacity, self.tokens + refill)
                    self.last_refill = now

                if self.tokens <= 0:
                    sleep_time = max(0.5, 60.0 - (now - self.last_refill))
                    time.sleep(sleep_time)
                    continue

                spacing = self.last_request + self.min_interval - now
                if spacing > 0:
                    time.sleep(spacing)
                    continue

                self.tokens -= 1
                self.last_request = time.time()
                break


@dataclass
class TextEntry:
    entry_id: str
    text: str
    location: str
    container: str
    index_path: tuple


class DocxBatchTranslator:
    def __init__(self, cfg: dict, glossary_buckets: dict, glossary_replace_terms: list, log_path: Path):
        api_key = cfg.get("openai_api_key") or os.getenv("OPENAI_API_KEY", "")
        if not api_key:
            raise RuntimeError("未设置 openai_api_key 且环境变量 OPENAI_API_KEY 为空")

        self.client = OpenAI(api_key=api_key, base_url=cfg.get("openai_base_url", "https://api.openai.com/v1"))
        self.model = cfg.get("model", "gpt-5.2")
        self.max_retries = max(int(cfg.get("max_retries", 3)), 1)
        self.prompt_profile = str(cfg.get("prompt_profile", "default")).strip().lower()
        self.request_timeout_seconds = float(cfg.get("request_timeout_seconds", 60))
        self.fallback_enabled = bool(cfg.get("fallback_on_failure_enabled", True))
        self.fallback_model = str(cfg.get("fallback_model_on_failure", "gpt-4.1")).strip()
        self.fallback_timeout_seconds = float(cfg.get("fallback_timeout_seconds", self.request_timeout_seconds))

        self.glossary_buckets = glossary_buckets
        self.glossary_replace_terms = glossary_replace_terms
        self.log_path = log_path

    def _build_instruction(self, glossary_matches):
        if self.prompt_profile == "dscryb":
            profile = (
                "你是资深TRPG场景文本翻译。"
                "将英文翻译为简体中文，保持沉浸感、画面感和自然中文节奏。"
            )
        else:
            profile = (
                "你是资深PF2E文本翻译。"
                "将英文翻译为简体中文，术语风格保持一致，规则词汇准确。"
            )

        glossary_hint = ""
        if glossary_matches:
            pairs = [f"{en} -> {cn}" for en, cn in glossary_matches[:80]]
            glossary_hint = "术语优先参考：" + "；".join(pairs) + "。"

        return (
            profile
            + glossary_hint
            + "你将收到一个段落数组。"
            "只翻译英文内容，保留原有段落含义。"
            "输出必须是严格JSON对象，且仅包含一个字段：translations。"
            "translations 必须是字符串数组，长度与输入段落数量完全一致，按原顺序对应。"
            "不要输出任何额外字段，不要解释，不要使用 markdown 代码块。"
        )

    def _call_model(self, model_name: str, timeout_seconds: float, instruction: str, payload: str):
        return self.client.responses.create(
            model=model_name,
            instructions=instruction,
            input=[
                {
                    "role": "user",
                    "content": [{"type": "input_text", "text": payload}],
                }
            ],
            timeout=timeout_seconds,
        )

    def _parse_translations(self, output_text: str, expected_len: int):
        obj = _parse_json_block(output_text)
        values = obj.get("translations") if isinstance(obj, dict) else None
        if not isinstance(values, list):
            return []

        parsed = []
        for item in values:
            if isinstance(item, str):
                parsed.append(_sanitize_text(item))
            elif isinstance(item, dict):
                parsed.append(_sanitize_text(str(item.get("text", ""))))
            else:
                parsed.append(_sanitize_text(str(item)))

        if len(parsed) != expected_len:
            return []
        return parsed

    def translate_batch(self, entries: list[TextEntry]):
        if not entries:
            return []

        all_text = "\n".join(e.text for e in entries)
        glossary_matches = _find_glossary_matches(all_text, self.glossary_buckets)
        instruction = self._build_instruction(glossary_matches)

        lines = []
        for idx, entry in enumerate(entries, start=1):
            lines.append(f"[{idx}] {entry.text}")

        payload = (
            f"请翻译以下 {len(entries)} 个段落，保持顺序：\n"
            + "\n".join(lines)
        )

        for attempt in range(1, self.max_retries + 1):
            try:
                resp = self._call_model(self.model, self.request_timeout_seconds, instruction, payload)
                output_text = (resp.output_text or "").strip()
                parsed = self._parse_translations(output_text, len(entries))
                if not parsed:
                    raise RuntimeError("模型返回格式不符合要求或数量不匹配")

                return [_apply_glossary_post_translation(t, self.glossary_replace_terms) for t in parsed]
            except Exception as e:
                if attempt < self.max_retries:
                    _write_log(self.log_path, f"RETRY batch attempt={attempt}/{self.max_retries} | {e}")
                    time.sleep(attempt)
                    continue

                if self.fallback_enabled and self.fallback_model:
                    try:
                        _write_log(self.log_path, f"FALLBACK model={self.fallback_model} | reason={e}")
                        resp = self._call_model(
                            self.fallback_model,
                            self.fallback_timeout_seconds,
                            instruction,
                            payload,
                        )
                        output_text = (resp.output_text or "").strip()
                        parsed = self._parse_translations(output_text, len(entries))
                        if parsed:
                            return [_apply_glossary_post_translation(t, self.glossary_replace_terms) for t in parsed]
                    except Exception as fb_err:
                        _write_log(self.log_path, f"ERROR fallback failed | {fb_err}")

                _write_log(self.log_path, f"ERROR batch translate failed | {e}")
                return []

        return []


def _iter_doc_entries(doc) -> list[TextEntry]:
    entries = []

    for p_idx, p in enumerate(doc.paragraphs):
        txt = _sanitize_text(p.text)
        entry_id = f"p:{p_idx}"
        entries.append(TextEntry(
            entry_id=entry_id,
            text=txt,
            location=f"paragraph[{p_idx}]",
            container="paragraph",
            index_path=(p_idx,),
        ))

    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for p_idx, p in enumerate(cell.paragraphs):
                    txt = _sanitize_text(p.text)
                    entry_id = f"t:{t_idx}:{r_idx}:{c_idx}:{p_idx}"
                    entries.append(TextEntry(
                        entry_id=entry_id,
                        text=txt,
                        location=f"table[{t_idx}].row[{r_idx}].cell[{c_idx}].paragraph[{p_idx}]",
                        container="table",
                        index_path=(t_idx, r_idx, c_idx, p_idx),
                    ))

    return entries


def _apply_translations_to_docx(doc, translated_map: dict, append_original: bool):
    for p_idx, p in enumerate(doc.paragraphs):
        entry_id = f"p:{p_idx}"
        if entry_id not in translated_map:
            continue
        translated = translated_map.get(entry_id, "")
        original = _sanitize_text(p.text)
        if not translated:
            continue
        p.text = f"{translated}\n{original}" if append_original and original else translated

    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for p_idx, p in enumerate(cell.paragraphs):
                    entry_id = f"t:{t_idx}:{r_idx}:{c_idx}:{p_idx}"
                    if entry_id not in translated_map:
                        continue
                    translated = translated_map.get(entry_id, "")
                    original = _sanitize_text(p.text)
                    if not translated:
                        continue
                    p.text = f"{translated}\n{original}" if append_original and original else translated


def _get_target_files(cfg: dict) -> list[Path]:
    folder_raw = str(cfg.get("input_folder_path", "")).strip()
    glob_pattern = str(cfg.get("input_file_glob", "*.docx")).strip() or "*.docx"
    recursive = bool(cfg.get("input_folder_recursive", False))

    files = []
    if folder_raw:
        folder = Path(folder_raw)
        if folder.exists() and folder.is_dir():
            iterator = folder.rglob(glob_pattern) if recursive else folder.glob(glob_pattern)
            files = sorted([p for p in iterator if p.is_file() and p.suffix.lower() == ".docx"])

    if files:
        return files

    single_path = Path(str(cfg.get("input_docx_path", "")).strip())
    if single_path and str(single_path) and single_path.exists() and single_path.is_file():
        return [single_path]
    return []


def _chunk_list(items: list, chunk_size: int):
    n = max(int(chunk_size), 1)
    return [items[i:i + n] for i in range(0, len(items), n)]


def _translate_one_docx(
    docx_path: Path,
    cfg: dict,
    translator: DocxBatchTranslator,
    limiter: RateLimiter,
    cache: dict,
    cache_enabled: bool,
    output_dir: Path,
    log_path: Path,
):
    source_doc = Document(str(docx_path))
    output_doc = Document(str(docx_path))

    entries = _iter_doc_entries(source_doc)
    candidates = []
    skip_count = 0

    for entry in entries:
        if not entry.text:
            skip_count += 1
            continue
        if _contains_chinese(entry.text):
            skip_count += 1
            continue
        if not _contains_english(entry.text):
            skip_count += 1
            continue
        candidates.append(entry)

    batch_size = max(int(cfg.get("batch_size", 20)), 1)
    batches = _chunk_list(candidates, batch_size)
    max_workers = max(int(cfg.get("max_workers", 8)), 1)

    stats = {"translated": 0, "cached": 0, "failed": 0, "skipped": skip_count}
    slowest = []
    translated_map = {}

    def worker(batch_entries: list[TextEntry]):
        if not batch_entries:
            return {}, 0, 0, 0, []

        local_map = {}
        local_cached = 0
        local_translated = 0
        local_failed = 0
        need_translate = []

        for entry in batch_entries:
            key = f"{docx_path.name}|{entry.entry_id}"
            text_hash = _hash_text(entry.text)
            cached = cache.get(key, {}) if cache_enabled else {}
            if cache_enabled and isinstance(cached, dict) and cached.get("hash") == text_hash:
                tr = _sanitize_text(str(cached.get("translated", "")))
                if tr:
                    local_map[entry.entry_id] = tr
                    local_cached += 1
                    continue
            need_translate.append(entry)

        if not need_translate:
            return local_map, local_cached, local_translated, local_failed, []

        limiter.wait()
        t0 = time.time()
        translated_list = translator.translate_batch(need_translate)
        elapsed = time.time() - t0

        if not translated_list or len(translated_list) != len(need_translate):
            local_failed += len(need_translate)
            return local_map, local_cached, local_translated, local_failed, [(elapsed, [e.entry_id for e in need_translate])]

        for entry, tr in zip(need_translate, translated_list):
            if tr:
                local_map[entry.entry_id] = tr
                local_translated += 1
                if cache_enabled:
                    key = f"{docx_path.name}|{entry.entry_id}"
                    with _cache_lock:
                        cache[key] = {
                            "hash": _hash_text(entry.text),
                            "translated": tr,
                            "updated_at": int(time.time()),
                        }
            else:
                local_failed += 1

        return local_map, local_cached, local_translated, local_failed, [(elapsed, [e.entry_id for e in need_translate])]

    desc = f"DOCX翻译 {docx_path.name}"
    with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as ex:
        futures = [ex.submit(worker, b) for b in batches]
        for fut in tqdm(concurrent.futures.as_completed(futures), total=len(futures), desc=desc, dynamic_ncols=True):
            try:
                local_map, c, t, f, slow_items = fut.result()
                translated_map.update(local_map)
                with _stats_lock:
                    stats["cached"] += c
                    stats["translated"] += t
                    stats["failed"] += f
                    slowest.extend(slow_items)
            except Exception as e:
                _write_log(log_path, f"ERROR batch future failed | file={docx_path.name} | {e}")

    append_original = bool(cfg.get("append_original_text", True))
    _apply_translations_to_docx(output_doc, translated_map, append_original)

    output_stem = docx_path.stem
    out_docx = output_dir / f"{output_stem}.translated.docx"
    output_doc.save(str(out_docx))

    ordered_rows = []
    for entry in entries:
        translated = translated_map.get(entry.entry_id, "")
        ordered_rows.append({
            "id": entry.entry_id,
            "location": entry.location,
            "source": entry.text,
            "translated": translated,
            "translated_or_cached": bool(translated),
        })

    if bool(cfg.get("output_json", True)):
        out_json = output_dir / f"{output_stem}.translated.json"
        out_json.write_text(json.dumps(ordered_rows, ensure_ascii=False, indent=2), encoding="utf-8")

    if bool(cfg.get("output_markdown", True)):
        out_md = output_dir / f"{output_stem}.translated.md"
        chunks = [f"# {output_stem} 翻译稿\n"]
        for row in ordered_rows:
            chunks.append(f"\n## {row['location']}\n")
            chunks.append((row["translated"] or "（未翻译/无需翻译）") + "\n")
            chunks.append("\n<details><summary>Source</summary>\n\n")
            chunks.append((row["source"] or "（空）") + "\n")
            chunks.append("\n</details>\n")
        out_md.write_text("".join(chunks), encoding="utf-8")

    top_n_slowest = max(int(cfg.get("top_n_slowest", 20)), 0)
    slowest.sort(key=lambda x: x[0], reverse=True)
    if top_n_slowest > 0 and slowest:
        print(f"最慢批次（{docx_path.name}）：")
        for sec, entry_ids in slowest[:top_n_slowest]:
            print(f"  {sec:.2f}s | entries={','.join(entry_ids[:6])}{'...' if len(entry_ids) > 6 else ''}")

    _save_cache(Path(cfg.get("cache_path", "docx_translate_cache.json")), cache, cache_enabled)

    print(
        f"完成 {docx_path.name}："
        f"翻译={stats['translated']} "
        f"缓存命中={stats['cached']} "
        f"跳过={stats['skipped']} "
        f"失败={stats['failed']}"
    )

    _write_log(
        log_path,
        f"DONE file={docx_path.name} translated={stats['translated']} cached={stats['cached']} "
        f"skipped={stats['skipped']} failed={stats['failed']}",
    )


def main():
    cfg = _load_config()
    log_path = Path(cfg.get("log_path", "docx_translate_run.log"))

    if Document is None:
        print("❌ 依赖缺失：请先安装 python-docx（pip install python-docx）")
        return

    targets = _get_target_files(cfg)
    if not targets:
        print(f"❌ 未找到待翻译 DOCX。请检查 {CONFIG_PATH} 的 input_docx_path 或 input_folder_path")
        return

    output_dir = Path(cfg.get("output_dir", "docx_output"))
    output_dir.mkdir(parents=True, exist_ok=True)

    cache_enabled = bool(cfg.get("cache_enabled", True))
    cache_path = Path(cfg.get("cache_path", "docx_translate_cache.json"))
    cache = _load_cache(cache_path, cache_enabled)

    glossary_path = Path(cfg.get("glossary_path", "glossary.json"))
    glossary_buckets, glossary_replace_terms, glossary_count = _load_glossary(glossary_path)

    print(f"目标DOCX：{len(targets)} 个")
    print(f"术语表加载：{glossary_count} 条")

    limiter = RateLimiter(int(cfg.get("target_rpm", 300)))
    translator = DocxBatchTranslator(cfg, glossary_buckets, glossary_replace_terms, log_path)

    _write_log(
        log_path,
        f"START files={len(targets)} glossary={glossary_count} cache={len(cache)} model={cfg.get('model', '')}",
    )

    for idx, docx_path in enumerate(targets, start=1):
        print(f"\n[{idx}/{len(targets)}] 处理：{docx_path}")
        try:
            _translate_one_docx(
                docx_path=docx_path,
                cfg=cfg,
                translator=translator,
                limiter=limiter,
                cache=cache,
                cache_enabled=cache_enabled,
                output_dir=output_dir,
                log_path=log_path,
            )
        except Exception as e:
            _write_log(log_path, f"ERROR file failed | file={docx_path} | {e}")
            print(f"❌ 处理失败：{docx_path} | {e}")

    _save_cache(cache_path, cache, cache_enabled)
    print("\n全部完成。")
    _write_log(log_path, "ALL_DONE")


if __name__ == "__main__":
    main()
