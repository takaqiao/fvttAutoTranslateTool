"""Extract Book 2 into formatting-aware segments.

Each paragraph becomes an ordered list of chunks. A chunk is either a plain run or a
hyperlink (which may itself span several runs). Translating chunk-by-chunk is what lets
us write the Chinese back into the docx without the link-displacement bug that mangled
Book 1 (where replacing runs[0] pushed hyperlink text to the end of the paragraph).

Consecutive runs sharing identical formatting are merged so the translator sees whole
phrases rather than Word's arbitrary run splits.
"""
import json
import os

from docx import Document
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'Season of Ghosts GM Guide - Book 2.docx')
DST = os.path.join(HERE, 'book2_segments.json')
IMG_DIR = os.path.join(HERE, 'book2_images')


# Body text is authored as near-black rather than "automatic"; treat it as unstyled.
DEFAULT_TEXT_COLORS = {'191813', '000000'}

# The author's semantic rating system lives in run shading (w:shd/@w:fill), not font colour.
SHADE_MEANING = {
    '6d9eeb': 'great',
    '93c47d': 'good',
    'ffd966': 'average',
    'e06666': 'bad',
    'b7b7b7': 'unrated',
    '8e7cc3': 'example',
    'f6b26b': 'extra-challenge',
    'c27ba0': 'legacy',
}


def run_fmt(r):
    """Formatting signature of a run, as plain JSON-able values."""
    color = None
    c = r.font.color
    if c is not None and c.type is not None and c.rgb is not None:
        color = str(c.rgb)
    if color and color.lower() in DEFAULT_TEXT_COLORS:
        color = None

    shade = None
    rpr = r._element.find(qn('w:rPr'))
    if rpr is not None:
        shd = rpr.find(qn('w:shd'))
        if shd is not None:
            fill = shd.get(qn('w:fill'))
            if fill and fill.lower() not in ('auto', 'ffffff'):
                shade = fill.lower()

    return {
        'bold': bool(r.bold),
        'italic': bool(r.italic),
        'strike': bool(r.font.strike),
        'color': color,
        'shade': shade,
    }


def run_text(r):
    parts = []
    for e in r._element.iter():
        if e.tag == qn('w:t'):
            if e.text:
                parts.append(e.text)
        elif e.tag == qn('w:tab'):
            parts.append('\t')
        elif e.tag == qn('w:br'):
            parts.append('\n')
    return ''.join(parts)


def list_level(p):
    """Return indent level for a numbered/bulleted paragraph, else None."""
    numpr = p._element.find('.//' + qn('w:numPr'))
    if numpr is None:
        return None
    ilvl = numpr.find(qn('w:ilvl'))
    return int(ilvl.get(qn('w:val'))) if ilvl is not None else 0


def para_chunks(p, rels):
    """Ordered chunks for a paragraph, walking w:r and w:hyperlink at the top level."""
    chunks = []
    for child in p._element:
        if child.tag == qn('w:r'):
            from docx.text.run import Run
            r = Run(child, p)
            txt = run_text(r)
            if not txt:
                continue
            fmt = run_fmt(r)
            if chunks and chunks[-1]['kind'] == 'text' and chunks[-1]['fmt'] == fmt:
                chunks[-1]['text'] += txt
            else:
                chunks.append({'kind': 'text', 'text': txt, 'fmt': fmt})
        elif child.tag == qn('w:hyperlink'):
            rid = child.get(qn('r:id'))
            href = rels[rid].target_ref if rid and rid in rels else None
            anchor = child.get(qn('w:anchor'))
            txt = ''.join(t.text or '' for t in child.iter(qn('w:t')))
            if not txt:
                continue
            first = child.find('.//' + qn('w:r'))
            fmt = {'bold': False, 'italic': False, 'strike': False,
                   'color': None, 'shade': None}
            if first is not None:
                from docx.text.run import Run
                fmt = run_fmt(Run(first, p))
            chunks.append({'kind': 'link', 'text': txt, 'fmt': fmt,
                           'href': href, 'anchor': anchor})
    return chunks


def export_images(doc):
    os.makedirs(IMG_DIR, exist_ok=True)
    mapping = {}  # paragraph index -> [filenames]
    n = 0
    for i, p in enumerate(doc.paragraphs):
        blips = p._element.findall('.//' + qn('a:blip'))
        for b in blips:
            rid = b.get(qn('r:embed'))
            if not rid:
                continue
            part = doc.part.related_parts[rid]
            n += 1
            ext = os.path.splitext(part.partname)[1] or '.png'
            fn = f'img{n:02d}{ext}'
            with open(os.path.join(IMG_DIR, fn), 'wb') as f:
                f.write(part.blob)
            mapping.setdefault(str(i), []).append(fn)
    return mapping


def main():
    doc = Document(SRC)
    rels = doc.part.rels

    segments = []
    for i, p in enumerate(doc.paragraphs):
        chunks = para_chunks(p, rels)
        text = ''.join(c['text'] for c in chunks)
        if not text.strip():
            continue
        segments.append({
            'i': i,
            'style': p.style.name if p.style else 'normal',
            'list': list_level(p),
            'chunks': chunks,
        })

    images = export_images(doc)

    with open(DST, 'w', encoding='utf-8') as f:
        json.dump({'segments': segments, 'images': images}, f,
                  ensure_ascii=False, indent=1)

    multi = sum(1 for s in segments if len(s['chunks']) > 1)
    print(f'segments: {len(segments)}  (multi-chunk: {multi})')
    print(f'chunks total: {sum(len(s["chunks"]) for s in segments)}')
    print(f'images: {sum(len(v) for v in images.values())} -> {IMG_DIR}')


if __name__ == '__main__':
    main()
